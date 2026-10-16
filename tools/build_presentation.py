from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


ROOT = Path(r"e:\PycharmProjects\TuristPRO")
MEDIA = ROOT / "presentation_assets" / "docx_media"
OUT = Path(
    r"C:\Users\user\Desktop\251\Создание системы бронирования экскурсий и мастер-классов в туристическом городе\TuristPRO_презентация.pptx"
)

GREEN = RGBColor(26, 107, 92)
DARK = RGBColor(30, 42, 40)
MUTED = RGBColor(89, 105, 101)
ACCENT = RGBColor(232, 168, 56)
LIGHT = RGBColor(244, 247, 246)
WHITE = RGBColor(255, 255, 255)


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_bg(slide, prs):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = LIGHT
    bg.line.fill.background()


def add_header(slide, prs, title: str, number: int | None = None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Cm(1.28))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()

    box = slide.shapes.add_textbox(Cm(0.7), Cm(0.26), prs.slide_width - Cm(2.0), Cm(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Segoe UI"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = WHITE

    if number is not None:
        num = slide.shapes.add_textbox(prs.slide_width - Cm(1.55), Cm(0.33), Cm(1.0), Cm(0.6))
        nt = num.text_frame
        nt.text = str(number)
        nt.paragraphs[0].alignment = PP_ALIGN.RIGHT
        nt.paragraphs[0].runs[0].font.name = "Segoe UI"
        nt.paragraphs[0].runs[0].font.size = Pt(13)
        nt.paragraphs[0].runs[0].font.color.rgb = WHITE


def add_footer(slide, prs):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Cm(0.7), prs.slide_height - Cm(0.42), prs.slide_width - Cm(1.4), Cm(0.035)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(216, 227, 224)
    line.line.fill.background()
    box = slide.shapes.add_textbox(Cm(0.7), prs.slide_height - Cm(0.38), prs.slide_width - Cm(1.4), Cm(0.25))
    tf = box.text_frame
    tf.text = "TuristPRO — система бронирования экскурсий и мастер-классов"
    r = tf.paragraphs[0].runs[0]
    r.font.name = "Segoe UI"
    r.font.size = Pt(7.5)
    r.font.color.rgb = MUTED


def add_title(slide, x, y, w, h, text, size=32, color=DARK, bold=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = "Segoe UI"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_body(slide, x, y, w, h, lines, size=16):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.space_after = Pt(7)
        p.font.name = "Segoe UI"
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
    return box


def add_card(slide, x, y, w, h, title, value):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(216, 227, 224)
    add_title(slide, x + Cm(0.35), y + Cm(0.25), w - Cm(0.7), Cm(0.6), title, size=12, color=MUTED, bold=True)
    add_title(slide, x + Cm(0.35), y + Cm(0.82), w - Cm(0.7), Cm(0.9), value, size=22, color=GREEN, bold=True)


def add_picture_fit(slide, image_path: Path, x, y, w, h):
    with Image.open(image_path) as im:
        iw, ih = im.size
    ratio = min(w / iw, h / ih)
    nw, nh = int(iw * ratio), int(ih * ratio)
    px = x + (w - nw) // 2
    py = y + (h - nh) // 2
    return slide.shapes.add_picture(str(image_path), px, py, width=nw, height=nh)


def add_image_panel(slide, image_name: str, x, y, w, h):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    panel.fill.solid()
    panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = RGBColor(216, 227, 224)
    add_picture_fit(slide, MEDIA / image_name, x + Cm(0.18), y + Cm(0.18), w - Cm(0.36), h - Cm(0.36))


def add_badge(slide, x, y, text):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Cm(3.4), Cm(0.55))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(232, 244, 241)
    shape.line.fill.background()
    tx = slide.shapes.add_textbox(x + Cm(0.18), y + Cm(0.08), Cm(3.0), Cm(0.35))
    tf = tx.text_frame
    tf.text = text
    r = tf.paragraphs[0].runs[0]
    r.font.name = "Segoe UI"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = GREEN


def build():
    prs = Presentation()
    prs.slide_width = Cm(33.867)
    prs.slide_height = Cm(19.05)

    # 1
    s = blank(prs)
    add_bg(s, prs)
    left = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Cm(10.2), prs.slide_height)
    left.fill.solid()
    left.fill.fore_color.rgb = GREEN
    left.line.fill.background()
    add_title(s, Cm(1.0), Cm(1.3), Cm(8.5), Cm(1.2), "TuristPRO", 38, WHITE)
    add_title(s, Cm(1.05), Cm(3.0), Cm(8.5), Cm(4.5), "Система бронирования экскурсий и мастер-классов в туристическом городе", 22, WHITE)
    add_title(s, Cm(11.2), Cm(3.5), Cm(20.5), Cm(4.0), "Дипломный проект", 34, DARK)
    add_body(s, Cm(11.25), Cm(7.6), Cm(18), Cm(2.4), ["Python · Flask · SQLAlchemy · SQLite", "Каталог, бронирование, оплата, личные кабинеты"], 18)
    add_footer(s, prs)

    # 2
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Актуальность проекта", 2)
    add_title(s, Cm(1), Cm(2.0), Cm(13.5), Cm(1), "Проблема", 28, GREEN)
    add_body(s, Cm(1), Cm(3.1), Cm(12.8), Cm(7), [
        "Запись на экскурсии часто ведётся через звонки и мессенджеры.",
        "Туристу сложно сравнивать даты, цены и свободные места.",
        "Организатору приходится вручную учитывать расписание и заявки.",
    ], 18)
    add_title(s, Cm(16.0), Cm(2.0), Cm(12.5), Cm(1), "Идея решения", 28, GREEN)
    add_body(s, Cm(16.0), Cm(3.1), Cm(12.8), Cm(7), [
        "Единый городской каталог мероприятий.",
        "Мгновенное бронирование с учётом лимита мест.",
        "Отдельные кабинеты туриста, организатора и администратора.",
    ], 18)
    add_footer(s, prs)

    # 3
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Цель и задачи", 3)
    add_title(s, Cm(1), Cm(2.1), Cm(31), Cm(1), "Цель", 24, GREEN)
    add_body(s, Cm(1), Cm(3.0), Cm(30), Cm(1.5), ["Разработать веб-приложение для бронирования экскурсий и мастер-классов в туристическом городе."], 19)
    add_card(s, Cm(1), Cm(5.3), Cm(7.5), Cm(2.1), "1", "Анализ")
    add_card(s, Cm(9.2), Cm(5.3), Cm(7.5), Cm(2.1), "2", "Проектирование")
    add_card(s, Cm(17.4), Cm(5.3), Cm(7.5), Cm(2.1), "3", "Реализация")
    add_card(s, Cm(25.6), Cm(5.3), Cm(7.2), Cm(2.1), "4", "Тестирование")
    add_body(s, Cm(1), Cm(9.0), Cm(30), Cm(3.2), [
        "MVP включает каталог, расписание, бронирование, имитацию оплаты и управление ролями.",
        "Результат: рабочий прототип TuristPRO с тестовыми данными и демонстрационными аккаунтами.",
    ], 18)
    add_footer(s, prs)

    # 4
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Архитектура приложения", 4)
    add_image_panel(s, "media_03.png", Cm(1), Cm(1.65), Cm(19.3), Cm(15.8))
    add_body(s, Cm(21.0), Cm(3.0), Cm(11.5), Cm(7.5), [
        "Клиент: браузер пользователя.",
        "Backend: Flask-приложение.",
        "Маршруты сгруппированы по blueprints.",
        "Данные хранятся через SQLAlchemy ORM.",
        "SQLite используется для дипломного прототипа.",
    ], 17)
    add_footer(s, prs)

    # 5
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Модель данных", 5)
    add_image_panel(s, "media_05.png", Cm(1), Cm(1.65), Cm(19.8), Cm(15.8))
    add_body(s, Cm(21.4), Cm(3.0), Cm(10.8), Cm(7.5), [
        "User — пользователь и роль.",
        "Activity — экскурсия или мастер-класс.",
        "Schedule — дата, время и места.",
        "Booking — заявка туриста.",
        "Payment — имитация оплаты.",
        "Category — раздел каталога.",
    ], 17)
    add_footer(s, prs)

    # 6
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Каталог мероприятий", 6)
    add_image_panel(s, "media_07.png", Cm(1), Cm(1.6), Cm(21.0), Cm(15.7))
    add_body(s, Cm(23.0), Cm(3.0), Cm(9.2), Cm(6), [
        "Карточки с изображениями.",
        "Категория, цена и длительность.",
        "Фильтрация по поиску, категории и цене.",
    ], 18)
    add_badge(s, Cm(23.0), Cm(9.7), "Экскурсии")
    add_badge(s, Cm(23.0), Cm(10.55), "Мастер-классы")
    add_badge(s, Cm(23.0), Cm(11.4), "Гастрономия")
    add_footer(s, prs)

    # 7
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Бронирование и оплата", 7)
    add_image_panel(s, "media_11.png", Cm(1.0), Cm(1.7), Cm(9.8), Cm(14.5))
    add_image_panel(s, "media_12.png", Cm(11.7), Cm(1.7), Cm(9.8), Cm(14.5))
    add_body(s, Cm(22.5), Cm(3.0), Cm(9.4), Cm(7), [
        "Проверяется количество свободных мест.",
        "После бронирования создаётся статус pending.",
        "Оплата имитируется без реального эквайринга.",
        "После успешной оплаты статус становится paid.",
    ], 16.5)
    add_footer(s, prs)

    # 8
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Личный кабинет туриста", 8)
    add_image_panel(s, "media_24.png", Cm(1), Cm(1.7), Cm(22.0), Cm(14.9))
    add_body(s, Cm(24.0), Cm(3.0), Cm(8.3), Cm(6.8), [
        "История бронирований.",
        "Дата и время мероприятия.",
        "Количество мест и сумма.",
        "Статус: pending, paid, cancelled.",
    ], 17)
    add_footer(s, prs)

    # 9
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Кабинет организатора", 9)
    add_image_panel(s, "media_17.png", Cm(1), Cm(1.7), Cm(14.7), Cm(14.9))
    add_image_panel(s, "media_18.png", Cm(16.2), Cm(1.7), Cm(16.0), Cm(14.9))
    add_footer(s, prs)

    # 10
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Администрирование", 10)
    add_image_panel(s, "media_19.png", Cm(1), Cm(1.7), Cm(15.4), Cm(14.9))
    add_image_panel(s, "media_23.png", Cm(17.0), Cm(1.7), Cm(15.2), Cm(14.9))
    add_body(s, Cm(1), Cm(16.7), Cm(31), Cm(0.7), ["Администратор контролирует пользователей, роли, категории, мероприятия и бронирования."], 14)
    add_footer(s, prs)

    # 11
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Тестирование", 11)
    add_image_panel(s, "media_27.png", Cm(1), Cm(1.7), Cm(14.7), Cm(14.9))
    add_image_panel(s, "media_28.png", Cm(16.2), Cm(1.7), Cm(16.0), Cm(14.9))
    add_footer(s, prs)

    # 12
    s = blank(prs)
    add_bg(s, prs)
    add_header(s, prs, "Итоги работы", 12)
    add_title(s, Cm(1.2), Cm(2.2), Cm(31), Cm(1), "Создан рабочий прототип TuristPRO", 30, GREEN)
    add_body(s, Cm(1.3), Cm(3.6), Cm(30), Cm(6.5), [
        "Система покрывает ключевой сценарий: каталог → выбор слота → бронирование → оплата → личный кабинет.",
        "Реализованы роли туриста, организатора и администратора.",
        "Прототип готов для демонстрации и может быть расширен реальной оплатой, отзывами и аналитикой.",
    ], 20)
    add_card(s, Cm(1.2), Cm(11.3), Cm(7.4), Cm(2.3), "12", "мероприятий")
    add_card(s, Cm(9.2), Cm(11.3), Cm(7.4), Cm(2.3), "48", "слотов")
    add_card(s, Cm(17.2), Cm(11.3), Cm(7.4), Cm(2.3), "3", "роли")
    add_card(s, Cm(25.2), Cm(11.3), Cm(7.4), Cm(2.3), "MVP", "готов")
    add_footer(s, prs)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
